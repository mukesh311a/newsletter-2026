import sys
from docx import Document
import os

doc = Document(r'c:\Users\Mukesh\Downloads\Newsletter 2026\Events.docx')

# Extract text
print("=== TEXT CONTENT ===")
for i, para in enumerate(doc.paragraphs):
    if para.text.strip():
        print(f"P{i}: [{para.style.name}] {para.text}")

# Extract tables
for ti, table in enumerate(doc.tables):
    print(f"\n=== TABLE {ti} ===")
    for ri, row in enumerate(table.rows):
        cells = [c.text.strip() for c in row.cells]
        print(f"R{ri}: {cells}")

# Extract images
out_dir = r'c:\Users\Mukesh\Downloads\Newsletter 2026\event_images'
os.makedirs(out_dir, exist_ok=True)
count = 0
for rel in doc.part.rels.values():
    if "image" in rel.reltype:
        count += 1
        img_data = rel.target_part.blob
        ext = rel.target_part.content_type.split('/')[-1].replace('jpeg','jpg')
        fname = f'event_{count}.{ext}'
        with open(os.path.join(out_dir, fname), 'wb') as f:
            f.write(img_data)
        print(f"Saved: {fname} ({len(img_data)} bytes)")

print(f"\nTotal images: {count}")
