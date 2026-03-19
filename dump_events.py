from docx import Document
doc = Document(r'c:\Users\Mukesh\Downloads\Newsletter 2026\Events.docx')
with open(r'c:\Users\Mukesh\Downloads\Newsletter 2026\events_text.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            f.write(f'P{i}: {p.text}\n')
    for ti, table in enumerate(doc.tables):
        f.write(f'\n=== TABLE {ti} ===\n')
        for ri, row in enumerate(table.rows):
            cells = [c.text.strip() for c in row.cells]
            f.write(f'R{ri}: {cells}\n')
print("Done")
